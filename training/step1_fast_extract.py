"""STEP 1: Ultra-fast streaming extractor for SCADA Excel files.
Aggregates 1s / 5min rows directly to 1-hour resolution per file, caches to parquet,
and produces the unified dataset in:
D:\\codes\\model files for agnitia hack it\\processed\\hourly_energy_2021_2025.parquet
"""
import io
import re
import sys
import time
import zipfile
import datetime
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor

import numpy as np
import pandas as pd
import openpyxl

ZIP_PATH = Path(r"D:\codes\datasets for agnitia hack it\Electricity Demand, Solar and Wind Generation Data.zip")
MODEL_DIR = Path(r"D:\codes\model files for agnitia hack it")
OUT_DIR = MODEL_DIR / "processed"
CACHE_DIR = OUT_DIR / "cache"
CACHE_DIR.mkdir(parents=True, exist_ok=True)


def extract_monthly_file(filename: str) -> str:
    """Extract and aggregate one monthly file to hourly resolution."""
    base = Path(filename).name
    cache_file = CACHE_DIR / (re.sub(r"[^A-Za-z0-9]+", "_", base) + ".parquet")
    if cache_file.exists():
        print(f"[cached] {base}", flush=True)
        return str(cache_file)

    t0 = time.time()
    print(f"[start] {base}", flush=True)
    with zipfile.ZipFile(ZIP_PATH) as zf:
        data = zf.read(filename)

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]

    hourly_acc = {}  # hour_ts -> [w_sum, s_sum, d_sum, count]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue
        t = row[0]
        if isinstance(t, datetime.datetime):
            hr = t.replace(minute=0, second=0, microsecond=0)
            w, s, d = row[5], row[6], row[10]
            if hr not in hourly_acc:
                hourly_acc[hr] = [0.0, 0.0, 0.0, 0]
            if w is not None and s is not None and d is not None:
                try:
                    hourly_acc[hr][0] += float(w)
                    hourly_acc[hr][1] += float(s)
                    hourly_acc[hr][2] += float(d)
                    hourly_acc[hr][3] += 1
                except (ValueError, TypeError):
                    pass
    wb.close()

    records = []
    for hr, (w_sum, s_sum, d_sum, cnt) in hourly_acc.items():
        if cnt > 0:
            records.append({
                "timestamp": hr,
                "wind_mw": w_sum / cnt,
                "solar_mw": s_sum / cnt,
                "demand_mw": d_sum / cnt,
            })

    df = pd.DataFrame(records)
    if not df.empty:
        df["timestamp"] = pd.to_datetime(df["timestamp"])
        df = df.sort_values("timestamp")
        df.to_parquet(cache_file, index=False)
        print(f"[done] {base} -> {len(df)} hours in {time.time()-t0:.1f}s", flush=True)
    return str(cache_file)


def extract_report_file(filename: str) -> str:
    """Extract January 2024 - June 2025 hourly report sheet."""
    base = Path(filename).name
    cache_file = CACHE_DIR / (re.sub(r"[^A-Za-z0-9]+", "_", base) + ".parquet")
    if cache_file.exists():
        print(f"[cached] {base}", flush=True)
        return str(cache_file)

    t0 = time.time()
    print(f"[start] {base}", flush=True)
    with zipfile.ZipFile(ZIP_PATH) as zf:
        data = zf.read(filename)

    df = pd.read_excel(io.BytesIO(data), sheet_name="Report")
    df.columns = [c.strip() for c in df.columns]
    df = df.rename(columns={
        "Timestamp": "timestamp",
        "Demand (MW)": "demand_mw",
        "Wind (MW)": "wind_mw",
        "Solar (MW)": "solar_mw",
    })
    df["timestamp"] = pd.to_datetime(df["timestamp"], format="%d-%m-%Y %H:%M:%S", errors="coerce")
    clean = df[["timestamp", "wind_mw", "solar_mw", "demand_mw"]].dropna(subset=["timestamp"])
    clean.to_parquet(cache_file, index=False)
    print(f"[done] {base} -> {len(clean)} hours in {time.time()-t0:.1f}s", flush=True)
    return str(cache_file)


def run_extraction():
    with zipfile.ZipFile(ZIP_PATH) as zf:
        all_files = sorted(n for n in zf.namelist() if n.lower().endswith(".xlsx"))

    print(f"Total Excel archives in zip: {len(all_files)}")
    cache_paths = []

    # Process all files
    for f in all_files:
        if "2024- June 2025" in f:
            cp = extract_report_file(f)
        else:
            cp = extract_monthly_file(f)
        cache_paths.append(cp)

    print("Consolidating all cached files into unified time-series...")
    frames = [pd.read_parquet(p) for p in cache_paths if Path(p).exists()]
    full = pd.concat(frames, ignore_index=True)
    full = full.drop_duplicates(subset=["timestamp"]).sort_values("timestamp").set_index("timestamp")

    # Complete hourly grid
    full_idx = pd.date_range(full.index.min(), full.index.max(), freq="1h")
    full = full.reindex(full_idx)
    full = full.interpolate(method="time", limit=24).bfill().ffill()

    # Clip negative values
    full["solar_mw"] = full["solar_mw"].clip(lower=0)
    full["wind_mw"] = full["wind_mw"].clip(lower=0)
    full["demand_mw"] = full["demand_mw"].clip(lower=0)

    full.index.name = "timestamp"
    final_path = OUT_DIR / "hourly_energy_2021_2025.parquet"
    full.to_parquet(final_path)
    print(f"SUCCESS! Saved unified dataset to: {final_path}")
    print(f"Total Rows: {len(full):,}, Date Range: {full.index.min()} -> {full.index.max()}")
    print("Columns summary:")
    print(full.describe().T[["count", "mean", "min", "max"]])


if __name__ == "__main__":
    run_extraction()
